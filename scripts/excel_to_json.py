"""
Converte a base original (Excel) em JSON para o site.

Uso:  python scripts/excel_to_json.py
Lê:   data/Base_4_BrightLearn_Cursos.xlsx   (abas: Dados, Dicionário de Dados, Como Usar)
Gera: src/data/matriculas.json   (dados BRUTOS, sem nenhum tratamento)
      src/data/dicionario.json   (dicionário + orientações de uso)

O tratamento/limpeza NÃO é feito aqui: ele acontece em src/lib/clean.ts,
para que o site consiga mostrar "antes e depois" de cada correção.
"""
import json
import pathlib

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parent.parent
XLSX = RAIZ / "data" / "Base_4_BrightLearn_Cursos.xlsx"
SAIDA = RAIZ / "src" / "data"


def limpa(v):
    """Mantém o valor bruto; só converte datas e números inteiros 'x.0'."""
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    if isinstance(v, float) and v.is_integer() and v < 1000:
        return int(v)
    return v


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # --- Aba Dados ---------------------------------------------------------
    ws = wb["Dados"]
    linhas = list(ws.iter_rows(values_only=True))
    cab = list(linhas[0])
    dados = [{c: limpa(v) for c, v in zip(cab, l)} for l in linhas[1:]]
    # Pct_Concluido e Nota_Final são decimais: mantém float
    for d in dados:
        for k in ("Pct_Concluido", "Nota_Final"):
            if d[k] is not None:
                d[k] = float(d[k])

    # --- Aba Dicionário de Dados ------------------------------------------
    ws = wb["Dicionário de Dados"]
    linhas = [tuple(c for c in l) for l in ws.iter_rows(values_only=True)]
    titulo = linhas[0][0]
    resumo = linhas[1][0]
    dic = [
        {"coluna": l[0], "tipo": l[1], "descricao": l[2]}
        for l in linhas[4:]
        if l[0]
    ]

    # --- Aba Como Usar ------------------------------------------------------
    ws = wb["Como Usar"]
    como = [
        str(l[0]).lstrip("• ").strip()
        for l in ws.iter_rows(values_only=True)
        if l[0] and str(l[0]).strip().startswith("•")
    ]

    SAIDA.mkdir(parents=True, exist_ok=True)
    (SAIDA / "matriculas.json").write_text(
        json.dumps(dados, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    (SAIDA / "dicionario.json").write_text(
        json.dumps(
            {"titulo": titulo, "resumo": resumo, "colunas": dic, "comoUsar": como},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"{len(dados)} matrículas, {len(dic)} colunas documentadas -> {SAIDA}")


if __name__ == "__main__":
    main()
